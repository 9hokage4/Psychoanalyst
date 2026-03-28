from openpyxl import load_workbook

wb = load_workbook('test_output.xlsx')
ws = wb['Группа']

print('=== Структура листа Группа ===')
for i in range(1, 20):
    row_values = [cell.value for cell in ws[i]]
    non_empty = [v for v in row_values if v]
    if non_empty:
        print(f'Row {i}: {non_empty[:6]}...')

print('\n=== Merged cells ===')
for merged_range in list(ws.merged_cells.ranges)[:10]:
    cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    print(f'Row {merged_range.min_row}: {cell.value}')
